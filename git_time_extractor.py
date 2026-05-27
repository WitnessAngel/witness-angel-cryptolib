"""
Estimate dev time spent on all WitnessAngel git repos for years 2023-2025.

Output: git_time_extractor.xlsx next to this script.
One sheet per year, one row per (deduplicated) commit:
    repo | commit_hash | description | author_email | hours

Time spent per commit = gap to previous commit by same author in same repo,
capped at 2.0h. First commit by an author gets 0.0h.

Branches: only the branch with the most recent tip commit per repo.
Dedup: within a repo, commits with identical subject + same changed paths
are collapsed to the earliest occurrence (rebase artifacts).
"""

from __future__ import annotations

import logging
import subprocess
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


WORKSPACE = Path(r"P:\Workspace\WitnessAngel")
SCRIPT_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = SCRIPT_DIR / "git_time_extractor.xlsx"
YEARS = (2023, 2024, 2025)
MAX_HOURS_PER_COMMIT = 2.0
GIT_FIELD_SEP = "\x1f"
GIT_RECORD_SEP = "\x1e"

log = logging.getLogger("git_time_extractor")


@dataclass
class Commit:
    sha: str
    timestamp: datetime
    email: str
    subject: str
    changed_paths: tuple = field(default_factory=tuple)
    hours: float = 0.0


def run_git(repo: Path, *args: str) -> str:
    result = subprocess.run(
        ["git", "-C", str(repo), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if result.returncode != 0:
        raise RuntimeError(f"git {args} failed in {repo.name}: {result.stderr.strip()}")
    return result.stdout


def find_repos(workspace: Path) -> list[Path]:
    return sorted(d for d in workspace.iterdir() if d.is_dir() and (d / ".git").is_dir())


def pick_latest_branch(repo: Path) -> str | None:
    out = run_git(
        repo,
        "for-each-ref",
        "--sort=-committerdate",
        "refs/heads/",
        "--format=%(refname:short)",
    )
    lines = [ln.strip() for ln in out.splitlines() if ln.strip()]
    return lines[0] if lines else None


def list_commits(repo: Path, branch: str) -> list[Commit]:
    fmt = GIT_FIELD_SEP.join(["%H", "%aI", "%ae", "%s"]) + GIT_RECORD_SEP
    out = run_git(repo, "log", branch, "--no-merges", f"--pretty=format:{fmt}")
    commits = []
    for rec in out.split(GIT_RECORD_SEP):
        rec = rec.strip("\n")
        if not rec:
            continue
        parts = rec.split(GIT_FIELD_SEP)
        if len(parts) != 4:
            continue
        sha, iso_ts, email, subject = parts
        try:
            ts = datetime.fromisoformat(iso_ts)
        except ValueError:
            continue
        commits.append(Commit(
            sha=sha,
            timestamp=ts,
            email=email or "<unknown>",
            subject=subject,
        ))
    return commits


def fetch_changed_paths(repo: Path, sha: str) -> tuple:
    out = run_git(repo, "show", "--name-only", "--pretty=format:", sha)
    paths = tuple(sorted(ln.strip() for ln in out.splitlines() if ln.strip()))
    return paths


def dedup_rebased(repo: Path, commits: list[Commit]) -> list[Commit]:
    # Only fetch changed_paths for commits with a duplicate subject — cheap optimization.
    subj_count = defaultdict(int)
    for c in commits:
        subj_count[c.subject] += 1
    suspects = {s for s, n in subj_count.items() if n > 1}

    if not suspects:
        return commits

    log.info("[%s] %d duplicate subjects — fetching changed paths", repo.name, len(suspects))
    for c in commits:
        if c.subject in suspects:
            try:
                c.changed_paths = fetch_changed_paths(repo, c.sha)
            except RuntimeError as e:
                log.warning("[%s] could not fetch paths for %s: %s", repo.name, c.sha[:8], e)

    seen = set()
    kept = []
    # Iterate in chronological order (oldest first) so we keep the earliest occurrence.
    for c in sorted(commits, key=lambda x: x.timestamp):
        if c.subject in suspects:
            key = (c.subject, c.changed_paths)
            if key in seen:
                continue
            seen.add(key)
        kept.append(c)
    return kept


def compute_hours(commits: list[Commit]) -> None:
    by_author: dict[str, list[Commit]] = defaultdict(list)
    for c in commits:
        by_author[c.email].append(c)
    for email, lst in by_author.items():
        lst.sort(key=lambda c: c.timestamp)
        prev_ts = None
        for c in lst:
            if prev_ts is None:
                c.hours = 0.0
            else:
                delta_h = (c.timestamp - prev_ts).total_seconds() / 3600.0
                c.hours = min(delta_h, MAX_HOURS_PER_COMMIT)
                if c.hours < 0:
                    c.hours = 0.0
            prev_ts = c.timestamp


def build_workbook() -> tuple[openpyxl.Workbook, dict]:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = {}
    for year in YEARS:
        sh = wb.create_sheet(str(year))
        sh.append(["repo", "commit_hash", "description", "author_email", "hours"])
        sheets[year] = sh
    return wb, sheets


def main() -> int:
    logging.basicConfig(
        stream=sys.stderr,
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%H:%M:%S",
    )

    log.info("Workspace: %s", WORKSPACE)
    log.info("Output:    %s", OUTPUT_PATH)

    if not WORKSPACE.is_dir():
        log.error("Workspace not found: %s", WORKSPACE)
        return 2

    repos = find_repos(WORKSPACE)
    log.info("Found %d git repos", len(repos))

    wb, sheets = build_workbook()

    for repo in repos:
        try:
            branch = pick_latest_branch(repo)
            if branch is None:
                log.warning("[%s] no local branches — skipping", repo.name)
                continue
            log.info("[%s] branch=%s", repo.name, branch)

            commits = list_commits(repo, branch)
            log.info("[%s] %d commits (pre-dedup)", repo.name, len(commits))

            commits = dedup_rebased(repo, commits)
            log.info("[%s] %d commits (post-dedup)", repo.name, len(commits))

            compute_hours(commits)

            written = defaultdict(int)
            for c in commits:
                year = c.timestamp.year
                if year in sheets:
                    sheets[year].append([
                        repo.name,
                        c.sha,
                        c.subject,
                        c.email,
                        round(c.hours, 3),
                    ])
                    written[year] += 1
            log.info(
                "[%s] wrote rows: %s",
                repo.name,
                ", ".join(f"{y}={written[y]}" for y in YEARS),
            )
        except Exception as e:
            log.warning("[%s] failed: %s", repo.name, e)
            continue

    for year, sh in sheets.items():
        last_data_row = sh.max_row
        if last_data_row >= 2:
            sh.append(["", "", "", "TOTAL", f"=SUM(E2:E{last_data_row})"])
        else:
            sh.append(["", "", "", "TOTAL", 0])
        log.info("Sheet %d: %d commit rows", year, max(0, last_data_row - 1))

    wb.save(OUTPUT_PATH)
    log.info("Wrote %s", OUTPUT_PATH)
    return 0


if __name__ == "__main__":
    sys.exit(main())
